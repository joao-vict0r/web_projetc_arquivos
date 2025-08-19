
import os
import mimetypes
import zipfile
import traceback
from flask import request, send_file, render_template
from werkzeug.utils import secure_filename
from pdf2docx import Converter
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from pypdf import PdfWriter, PdfReader
import ffmpeg
import pandas as pd
from fpdf import FPDF
from io import BytesIO
from app.utils.cleanup import cleanup_files

UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'uploads')

def handle_conversion(request):
	action = request.form.get('action')
	if action in ('pdf2word', 'word2pdf', 'video2mp3', 'unzipfile', 'excel2pdf'):
		files = request.files.getlist('file')
	elif action == 'zipfile':
		files = request.files.getlist('zipfiles')
	else:
		return 'Ação inválida ou não suportada.', 400
	if not files or files[0].filename == '':
		return 'Nenhum arquivo enviado.', 400
	filepaths = []
	filenames = []
	allowed_exts = {'.pdf', '.docx', '.zip', '.mp4', '.png', '.jpg', '.jpeg', '.txt', '.csv', '.xls', '.xlsx', '.ppt', '.pptx', '.gif', '.bmp', '.rtf', '.odt', '.ods', '.odp', '.svg', '.webp', '.mp3', '.wav', '.avi', '.mov', '.wmv', '.flv', '.mkv', '.json', '.xml', '.html', '.htm', '.md', '.yml', '.yaml', '.ini', '.log'}
	for file in files:
		filename = secure_filename(file.filename)
		ext = os.path.splitext(filename)[1].lower()
		if ext not in allowed_exts:
			return f'Extensão de arquivo não permitida: {ext}', 400
		filepath = os.path.join(UPLOAD_FOLDER, filename)
		if not os.path.abspath(filepath).startswith(os.path.abspath(UPLOAD_FOLDER)):
			return 'Caminho de arquivo inválido.', 400
		file.save(filepath)
		filepaths.append(filepath)
		filenames.append(filename)
	if not filenames:
		return 'Nenhum arquivo enviado.', 400
	name_without_ext = os.path.splitext(filenames[0])[0]

	# ==============================
	# DESCOMPACTAR ZIP
	# ==============================
	if action == 'unzipfile':
		zip_path = filepaths[0]
		if not zipfile.is_zipfile(zip_path):
			os.remove(zip_path)
			return 'O arquivo enviado não é um ZIP válido.', 400

		extract_dir = os.path.join(UPLOAD_FOLDER, f"extract_{name_without_ext}")
		os.makedirs(extract_dir, exist_ok=True)
		try:
			with zipfile.ZipFile(zip_path, 'r') as zip_ref:
				zip_ref.extractall(extract_dir)
			os.remove(zip_path)
			extracted_files = []
			for root, _, files_in_dir in os.walk(extract_dir):
				for file_in_dir in files_in_dir:
					abs_path = os.path.join(root, file_in_dir)
					rel_path = os.path.relpath(abs_path, extract_dir)
					extracted_files.append(rel_path)
			return render_template('unzip_list.html', files=extracted_files, extract_dir=os.path.basename(extract_dir))
		except Exception as e:
			tb = traceback.format_exc()
			return f'Erro ao descompactar: {str(e)}\n{tb}', 500

	elif action == 'zipfile':
		zip_path = os.path.join(UPLOAD_FOLDER, f"{secure_filename(name_without_ext)}.zip")
		try:
			with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
				for fp, fn in zip(filepaths, filenames):
					zipf.write(fp, arcname=secure_filename(fn))
			cleanup_files(filepaths + [zip_path])
			mime = mimetypes.guess_type(zip_path)[0] or 'application/zip'
			return send_file(zip_path, as_attachment=True, download_name=os.path.basename(zip_path), mimetype=mime)
		except Exception as e:
			return f'Erro ao compactar: {str(e)}', 500

	elif action == 'pdf2word' and filenames[0].lower().endswith('.pdf'):
		docx_path = os.path.join(UPLOAD_FOLDER, f"{secure_filename(name_without_ext)}.docx")
		try:
			cv = Converter(filepaths[0])
			cv.convert(docx_path, start=0, end=None)
			cv.close()
			cleanup_files([filepaths[0], docx_path])
			mime = mimetypes.guess_type(docx_path)[0] or 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
			return send_file(docx_path, as_attachment=True, download_name=os.path.basename(docx_path), mimetype=mime)
		except Exception as e:
			return f'Erro ao converter PDF: {str(e)}', 500

	elif action == 'word2pdf' and filenames[0].lower().endswith('.docx'):
		pdf_path = os.path.join(UPLOAD_FOLDER, f"{secure_filename(name_without_ext)}.pdf")
		try:
			doc = Document(filepaths[0])
			pdf_writer = PdfWriter()
			for para in doc.paragraphs:
				packet = BytesIO()
				can = canvas.Canvas(packet, pagesize=letter)
				can.drawString(100, 750, para.text)
				can.save()
				packet.seek(0)
				pdf_reader = PdfReader(packet)
				pdf_writer.add_page(pdf_reader.pages[0])
			with open(pdf_path, 'wb') as f:
				pdf_writer.write(f)
			cleanup_files([filepaths[0], pdf_path])
			mime = mimetypes.guess_type(pdf_path)[0] or 'application/pdf'
			return send_file(pdf_path, as_attachment=True, download_name=os.path.basename(pdf_path), mimetype=mime)
		except Exception as e:
			return f'Erro ao converter Word: {str(e)}', 500

	elif action == 'excel2pdf' and filenames[0].lower().endswith(('.xls', '.xlsx', '.csv', '.ods')):
		pdf_path = os.path.join(UPLOAD_FOLDER, f"{secure_filename(name_without_ext)}.pdf")
		try:
			ext = os.path.splitext(filenames[0])[1].lower()
			if ext == '.csv':
				df = pd.read_csv(filepaths[0])
			elif ext in ('.xls', '.xlsx'):
				df = pd.read_excel(filepaths[0])
			elif ext == '.ods':
				df = pd.read_excel(filepaths[0], engine='odf')
			else:
				return 'Formato não suportado para conversão.', 400

			pdf = FPDF()
			pdf.add_page()
			pdf.set_font('Arial', size=10)
			col_width = pdf.w / (len(df.columns) + 1)
			row_height = pdf.font_size * 1.5
			for col in df.columns:
				pdf.cell(col_width, row_height, str(col), border=1)
			pdf.ln(row_height)
			for i, row in df.iterrows():
				for item in row:
					pdf.cell(col_width, row_height, str(item), border=1)
				pdf.ln(row_height)
			pdf.output(pdf_path)
			cleanup_files([filepaths[0], pdf_path])
			mime = mimetypes.guess_type(pdf_path)[0] or 'application/pdf'
			return send_file(pdf_path, as_attachment=True, download_name=os.path.basename(pdf_path), mimetype=mime)
		except Exception as e:
			return f'Erro ao converter planilha: {str(e)}', 500

	elif action == 'video2mp3':
		video_exts = ['.mp4', '.avi', '.mkv', '.mov', '.wmv', '.flv', '.webm', '.3gp', '.mpeg', '.mpg', '.ogv']
		ext = os.path.splitext(filenames[0])[1].lower()
		if ext not in video_exts:
			return f'Formato de vídeo não suportado: {ext}', 400
		mp3_path = os.path.join(UPLOAD_FOLDER, f"{name_without_ext}.mp3")
		try:
			(
				ffmpeg
				.input(filepaths[0])
				.output(mp3_path, format='mp3', acodec='libmp3lame', audio_bitrate='128k')
				.run(overwrite_output=True)
			)
			cleanup_files([filepaths[0], mp3_path])
			return send_file(mp3_path, as_attachment=True)
		except Exception as e:
			return f'Erro ao converter vídeo: {str(e)}', 500

	return 'Arquivo ou ação inválida. Envie PDF, DOCX, MP4 ou selecione ZIP.', 400
